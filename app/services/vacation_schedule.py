"""Projected vacation schedule for customer print / Excel handoff."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.models import Employee, LeaveRequest, LeaveStatus, LeaveType, Role
from app.services.fosc_export import quarter_bounds
from app.services.leave_balance import count_workdays, get_leave_balance, iter_workdays


CONTRACT_TITLE = "Follow-On Support Contract (FOSC)"
DOCUMENT_TITLE = "Projected Vacation Schedule"
PROGRAM_LINE = "SDC In-Country Team"


@dataclass(frozen=True)
class Period:
    year: int
    quarter: int | None
    start: date
    end: date

    @property
    def label(self) -> str:
        if self.quarter:
            return f"Q{self.quarter} {self.year}"
        return f"Calendar Year {self.year}"

    @property
    def span_label(self) -> str:
        return f"{self.start.strftime('%d %b %Y')} – {self.end.strftime('%d %b %Y')}"

    @property
    def days(self) -> int:
        return (self.end - self.start).days + 1


def resolve_period(year: int, quarter: int | None) -> Period:
    if quarter in (1, 2, 3, 4):
        start, end = quarter_bounds(year, quarter)
        return Period(year=year, quarter=quarter, start=start, end=end)
    return Period(year=year, quarter=None, start=date(year, 1, 1), end=date(year, 12, 31))


def _clip(start: date, end: date, period: Period) -> tuple[date, date] | None:
    s = max(start, period.start)
    e = min(end, period.end)
    if e < s:
        return None
    return s, e


def _pct(day: date, period: Period, *, inclusive_end: bool = False) -> float:
    """Position of a date as percent of the period width."""
    span = period.days
    if span <= 0:
        return 0.0
    offset = (day - period.start).days
    if inclusive_end:
        offset += 1
    return max(0.0, min(100.0, offset / span * 100.0))


def _month_headers(period: Period) -> list[dict]:
    months: list[dict] = []
    cursor = date(period.start.year, period.start.month, 1)
    while cursor <= period.end:
        if cursor.month == 12:
            month_end = date(cursor.year, 12, 31)
        else:
            month_end = date(cursor.year, cursor.month + 1, 1) - timedelta(days=1)
        vis_start = max(cursor, period.start)
        vis_end = min(month_end, period.end)
        left = _pct(vis_start, period)
        right = _pct(vis_end, period, inclusive_end=True)
        months.append({
            "key": cursor.strftime("%Y-%m"),
            "label": cursor.strftime("%b"),
            "year": cursor.year,
            "left": round(left, 3),
            "width": round(max(0.0, right - left), 3),
        })
        if cursor.month == 12:
            cursor = date(cursor.year + 1, 1, 1)
        else:
            cursor = date(cursor.year, cursor.month + 1, 1)
    return months


def _types(include_holidays: bool) -> list[LeaveType]:
    types = [LeaveType.vacation]
    if include_holidays:
        types.append(LeaveType.uae_holiday)
    return types


def _statuses(include_pending: bool) -> list[LeaveStatus]:
    statuses = [LeaveStatus.approved]
    if include_pending:
        statuses.append(LeaveStatus.pending)
    return statuses


def build_vacation_schedule(
    db: Session,
    year: int,
    quarter: int | None = None,
    *,
    include_pending: bool = False,
    include_holidays: bool = True,
    prepared_by: str | None = None,
    as_of: date | None = None,
) -> dict:
    """Team projected vacation for a year or quarter, ready for print/Excel."""
    period = resolve_period(year, quarter)
    as_of = as_of or date.today()
    employees = (
        db.query(Employee)
        .filter(Employee.is_active.is_(True))
        .order_by(Employee.name)
        .all()
    )

    requests = (
        db.query(LeaveRequest)
        .options(joinedload(LeaveRequest.employee))
        .filter(
            LeaveRequest.leave_type.in_(_types(include_holidays)),
            LeaveRequest.status.in_(_statuses(include_pending)),
            LeaveRequest.start_date <= period.end,
            LeaveRequest.end_date >= period.start,
        )
        .order_by(LeaveRequest.start_date, LeaveRequest.end_date)
        .all()
    )

    trips: list[dict] = []
    for req in requests:
        emp = req.employee
        if emp is None or not emp.is_active:
            continue
        clipped = _clip(req.start_date, req.end_date, period)
        if clipped is None:
            continue
        c_start, c_end = clipped
        workdays = count_workdays(c_start, c_end)
        left = _pct(c_start, period)
        right = _pct(c_end, period, inclusive_end=True)
        trips.append({
            "leave_id": req.id,
            "employee_id": emp.id,
            "employee_name": emp.name,
            "role": emp.role.value if emp.role else Role.employee.value,
            "leave_type": req.leave_type.value,
            "status": req.status.value,
            "start": req.start_date,
            "end": req.end_date,
            "clipped_start": c_start,
            "clipped_end": c_end,
            "workdays": workdays,
            "comments": (req.comments or "").strip(),
            "left": round(left, 3),
            "width": round(max(0.4, right - left), 3),
            "bar_label": _bar_label(c_start, c_end),
        })

    trips.sort(key=lambda t: (t["employee_name"].lower(), t["clipped_start"], t["leave_id"]))

    occupancy: dict[date, list[str]] = {}
    for trip in trips:
        for d in iter_workdays(trip["clipped_start"], trip["clipped_end"]):
            occupancy.setdefault(d, []).append(trip["employee_name"])

    peak_count = max((len(v) for v in occupancy.values()), default=0)
    peak_days = sorted(d for d, names in occupancy.items() if len(names) == peak_count) if peak_count else []
    overlap_days = sorted(d for d, names in occupancy.items() if len(names) >= 2)

    roster: list[dict] = []
    for emp in employees:
        emp_trips = [t for t in trips if t["employee_id"] == emp.id]
        vac_trips = [t for t in emp_trips if t["leave_type"] == LeaveType.vacation.value]
        hol_trips = [t for t in emp_trips if t["leave_type"] == LeaveType.uae_holiday.value]
        balance = get_leave_balance(db, emp, year=year)
        roster.append({
            "employee_id": emp.id,
            "name": emp.name,
            "role": emp.role.value if emp.role else Role.employee.value,
            "vacation_workdays": sum(t["workdays"] for t in vac_trips),
            "holiday_workdays": sum(t["workdays"] for t in hol_trips),
            "trip_count": len(emp_trips),
            "trips": emp_trips,
            "has_pending": any(t["status"] == LeaveStatus.pending.value for t in emp_trips),
            "vacation_remaining": int(balance["vacation"]["remaining"]),
            "vacation_entitlement": int(balance["vacation"]["entitlement"]),
            "vacation_used": int(balance["vacation"]["used"]),
            "vacation_pending": int(balance["vacation"]["pending"]),
        })

    gantt_rows = [r for r in roster if r["trips"]]
    none_scheduled = [r["name"] for r in roster if not r["trips"]]

    occupancy_rows = [
        {
            "date": d,
            "weekday": d.strftime("%a"),
            "count": len(names),
            "names": names,
            "on_duty": len(employees) - len(set(names)),
        }
        for d, names in sorted(occupancy.items())
    ]

    return {
        "title": DOCUMENT_TITLE,
        "contract": CONTRACT_TITLE,
        "program": PROGRAM_LINE,
        "period": period,
        "year": year,
        "quarter": period.quarter,
        "as_of": as_of,
        "prepared_by": prepared_by or "",
        "include_pending": include_pending,
        "include_holidays": include_holidays,
        "staff_count": len(employees),
        "trip_count": len(trips),
        "vacation_trip_count": sum(1 for t in trips if t["leave_type"] == LeaveType.vacation.value),
        "pending_count": sum(1 for t in trips if t["status"] == LeaveStatus.pending.value),
        "workdays_on_leave": len(occupancy),
        "peak_count": peak_count,
        "peak_days": peak_days[:8],
        "overlap_days": overlap_days,
        "overlap_count": len(overlap_days),
        "months": _month_headers(period),
        "trips": trips,
        "roster": roster,
        "gantt_rows": gantt_rows,
        "none_scheduled": none_scheduled,
        "occupancy_rows": occupancy_rows,
        "occupancy_hot": [row for row in occupancy_rows if row["count"] >= 2],
    }


def _bar_label(start: date, end: date) -> str:
    if start == end:
        return start.strftime("%d %b")
    if start.month == end.month and start.year == end.year:
        return f"{start.day}–{end.strftime('%d %b')}"
    return f"{start.strftime('%d %b')}–{end.strftime('%d %b')}"


def format_date_list(days: Iterable[date], limit: int = 6) -> str:
    days = list(days)
    if not days:
        return "None"
    shown = [d.strftime("%d %b") for d in days[:limit]]
    extra = len(days) - limit
    if extra > 0:
        shown.append(f"+{extra} more")
    return ", ".join(shown)


def build_vacation_schedule_workbook(schedule: dict) -> BytesIO:
    """Customer Excel: roster trips + coverage days."""
    wb = Workbook()
    navy = PatternFill("solid", fgColor="1B2A4A")
    navy_font = Font(bold=True, color="FFFFFF", name="Calibri")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    title_font = Font(bold=True, color="FFFFFF", name="Calibri", size=16)
    thin = Border(
        left=Side(style="thin", color="C5C0B5"),
        right=Side(style="thin", color="C5C0B5"),
        top=Side(style="thin", color="C5C0B5"),
        bottom=Side(style="thin", color="C5C0B5"),
    )
    pending_fill = PatternFill("solid", fgColor="F4E6C4")
    holiday_fill = PatternFill("solid", fgColor="D9E6F2")
    hot_fill = PatternFill("solid", fgColor="F3D6D0")
    alt_fill = PatternFill("solid", fgColor="F4F1EA")

    period: Period = schedule["period"]

    cover = wb.active
    cover.title = "Cover"
    cover.merge_cells("A1:F1")
    cover["A1"] = DOCUMENT_TITLE
    cover["A1"].font = title_font
    cover["A1"].fill = navy
    cover["A1"].alignment = Alignment(horizontal="left", vertical="center")
    cover.row_dimensions[1].height = 28
    cover["A3"] = CONTRACT_TITLE
    cover["A4"] = PROGRAM_LINE
    cover["A5"] = f"Period: {period.label}  ({period.span_label})"
    cover["A6"] = f"Issued: {schedule['as_of'].strftime('%d %b %Y')}"
    if schedule.get("prepared_by"):
        cover["A7"] = f"Prepared by: {schedule['prepared_by']}"
    cover["A9"] = f"Active staff: {schedule['staff_count']}"
    cover["A10"] = f"Vacation periods in window: {schedule['vacation_trip_count']}"
    cover["A11"] = f"Peak concurrent leave: {schedule['peak_count']} staff"
    cover["A12"] = f"Peak dates: {format_date_list(schedule['peak_days'])}"
    cover["A13"] = f"Days with 2+ staff on leave: {schedule['overlap_count']}"
    if schedule["include_pending"]:
        cover["A15"] = "Includes pending requests (not yet approved)."
    else:
        cover["A15"] = "Approved leave only (customer copy)."
    cover["A17"] = "Prepared by (SDC):"
    cover["A18"] = "Acknowledged by (Customer):"
    cover.column_dimensions["A"].width = 55

    ws = wb.create_sheet("Roster")
    headers = [
        "Employee", "Role", "Type", "Status", "Start", "End",
        "Workdays", "Comments", "Vacation remaining (year)",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.font = header_font
        cell.fill = navy
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[1].height = 22

    row = 2
    for emp in schedule["roster"]:
        if not emp["trips"]:
            values = [
                emp["name"], emp["role"], "—", "—", "—", "—", 0, "No projected vacation",
                emp["vacation_remaining"],
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row, col, val)
                cell.border = thin
                if row % 2 == 0:
                    cell.fill = alt_fill
            row += 1
            continue
        for trip in emp["trips"]:
            values = [
                emp["name"],
                emp["role"],
                trip["leave_type"].replace("_", " ").title(),
                trip["status"].title(),
                trip["clipped_start"].isoformat(),
                trip["clipped_end"].isoformat(),
                trip["workdays"],
                trip["comments"],
                emp["vacation_remaining"],
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row, col, val)
                cell.border = thin
                if trip["status"] == "pending":
                    cell.fill = pending_fill
                elif trip["leave_type"] == "uae_holiday":
                    cell.fill = holiday_fill
                elif row % 2 == 0:
                    cell.fill = alt_fill
            row += 1

    widths = [22, 14, 16, 12, 14, 14, 12, 36, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:I{max(1, row - 1)}"
    ws.freeze_panes = "A2"

    cov = wb.create_sheet("Coverage")
    cov_headers = ["Date", "Weekday", "Staff on leave", "Names", "Staff on duty"]
    for i, h in enumerate(cov_headers, 1):
        cell = cov.cell(1, i, h)
        cell.font = header_font
        cell.fill = navy
        cell.border = thin
    for i, occ in enumerate(schedule["occupancy_rows"], 2):
        vals = [
            occ["date"].isoformat(),
            occ["weekday"],
            occ["count"],
            ", ".join(occ["names"]),
            occ["on_duty"],
        ]
        for col, val in enumerate(vals, 1):
            cell = cov.cell(i, col, val)
            cell.border = thin
            if occ["count"] >= 2:
                cell.fill = hot_fill
    for i, w in enumerate([14, 12, 16, 50, 16], 1):
        cov.column_dimensions[get_column_letter(i)].width = w
    cov.freeze_panes = "A2"
    cov.auto_filter.ref = f"A1:E{max(1, 1 + len(schedule['occupancy_rows']))}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
