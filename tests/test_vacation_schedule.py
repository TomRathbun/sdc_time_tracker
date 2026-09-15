"""Projected vacation schedule: period clip, occupancy, customer Excel."""

from __future__ import annotations

import unittest
from datetime import date

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.auth import hash_pin
from app.database import Base
from app.models import Employee, LeaveRequest, LeaveStatus, LeaveType, Role
from app.services.vacation_schedule import (
    build_vacation_schedule,
    build_vacation_schedule_workbook,
    resolve_period,
)


class VacationScheduleTests(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        self.tom = self._emp("Tom Rathbun", Role.manager)
        self.omar = self._emp("Omar Eldeeb", Role.supervisor)
        self.mano = self._emp("Manoj Nada", Role.employee)
        self.inactive = self._emp("Inactive Person", Role.employee, active=False)

    def tearDown(self):
        self.db.close()

    def _emp(self, name, role, active=True):
        emp = Employee(
            name=name,
            pin_hash=hash_pin("1234"),
            role=role,
            is_active=active,
            pin_needs_reset=False,
        )
        self.db.add(emp)
        self.db.commit()
        self.db.refresh(emp)
        return emp

    def _leave(self, emp, start, end, status=LeaveStatus.approved, kind=LeaveType.vacation, comments=""):
        req = LeaveRequest(
            employee_id=emp.id,
            leave_type=kind,
            start_date=start,
            end_date=end,
            status=status,
            comments=comments,
        )
        self.db.add(req)
        self.db.commit()
        self.db.refresh(req)
        return req

    def test_quarter_bounds(self):
        p = resolve_period(2026, 4)
        self.assertEqual(p.start, date(2026, 10, 1))
        self.assertEqual(p.end, date(2026, 12, 31))
        self.assertEqual(p.label, "Q4 2026")

    def test_year_default(self):
        p = resolve_period(2026, None)
        self.assertEqual(p.start, date(2026, 1, 1))
        self.assertEqual(p.end, date(2026, 12, 31))

    def test_clips_to_quarter_and_counts_workdays(self):
        # Mon 28 Sep – Fri 2 Oct 2026 spans Q3/Q4. Q4 clip is Thu 1 Oct + Fri 2 Oct = 2 workdays.
        self._leave(self.tom, date(2026, 9, 28), date(2026, 10, 2), comments="US trip")
        sched = build_vacation_schedule(self.db, 2026, 4, include_pending=False, include_holidays=False)
        trips = [t for t in sched["trips"] if t["employee_name"] == "Tom Rathbun"]
        self.assertEqual(len(trips), 1)
        self.assertEqual(trips[0]["clipped_start"], date(2026, 10, 1))
        self.assertEqual(trips[0]["clipped_end"], date(2026, 10, 2))
        self.assertEqual(trips[0]["workdays"], 2)

    def test_pending_excluded_by_default(self):
        self._leave(self.omar, date(2026, 11, 8), date(2026, 11, 12), status=LeaveStatus.pending)
        hidden = build_vacation_schedule(self.db, 2026, 4, include_pending=False)
        self.assertEqual(hidden["trip_count"], 0)
        shown = build_vacation_schedule(self.db, 2026, 4, include_pending=True)
        self.assertEqual(shown["trip_count"], 1)
        self.assertEqual(shown["pending_count"], 1)
        self.assertTrue(shown["gantt_rows"][0]["has_pending"])

    def test_weekend_only_range_has_zero_workdays(self):
        self._leave(self.mano, date(2026, 10, 3), date(2026, 10, 4))  # Sat–Sun
        sched = build_vacation_schedule(self.db, 2026, 4, include_holidays=False)
        self.assertEqual(sched["trips"][0]["workdays"], 0)
        self.assertEqual(sched["workdays_on_leave"], 0)

    def test_occupancy_and_peak(self):
        self._leave(self.tom, date(2026, 10, 12), date(2026, 10, 16))   # Mon–Fri
        self._leave(self.omar, date(2026, 10, 14), date(2026, 10, 15))  # Wed–Thu overlap
        sched = build_vacation_schedule(self.db, 2026, 4, include_holidays=False)
        self.assertEqual(sched["peak_count"], 2)
        self.assertIn(date(2026, 10, 14), sched["peak_days"])
        self.assertIn(date(2026, 10, 15), sched["peak_days"])
        self.assertEqual(sched["overlap_count"], 2)
        self.assertEqual(sched["staff_count"], 3)  # inactive omitted

    def test_inactive_staff_omitted(self):
        self._leave(self.inactive, date(2026, 10, 5), date(2026, 10, 9))
        sched = build_vacation_schedule(self.db, 2026, 4)
        names = [r["name"] for r in sched["roster"]]
        self.assertNotIn("Inactive Person", names)

    def test_holidays_optional(self):
        self._leave(
            self.mano,
            date(2026, 12, 2),
            date(2026, 12, 3),
            kind=LeaveType.uae_holiday,
            comments="UAE National Day",
        )
        without = build_vacation_schedule(self.db, 2026, 4, include_holidays=False)
        self.assertEqual(without["trip_count"], 0)
        with_h = build_vacation_schedule(self.db, 2026, 4, include_holidays=True)
        self.assertEqual(with_h["trip_count"], 1)
        self.assertEqual(with_h["trips"][0]["leave_type"], "uae_holiday")

    def test_none_scheduled_listed(self):
        self._leave(self.tom, date(2026, 12, 20), date(2026, 12, 24))
        sched = build_vacation_schedule(self.db, 2026, 4, include_holidays=False)
        self.assertIn("Omar Eldeeb", sched["none_scheduled"])
        self.assertIn("Manoj Nada", sched["none_scheduled"])
        self.assertNotIn("Tom Rathbun", sched["none_scheduled"])

    def test_excel_workbook_has_three_sheets(self):
        self._leave(self.tom, date(2026, 12, 20), date(2026, 12, 24), comments="Year-end")
        sched = build_vacation_schedule(
            self.db, 2026, 4, include_holidays=False, prepared_by="Tom Rathbun"
        )
        buf = build_vacation_schedule_workbook(sched)
        self.assertGreater(len(buf.getvalue()), 1000)
        from openpyxl import load_workbook
        wb = load_workbook(buf)
        self.assertEqual(wb.sheetnames, ["Cover", "Roster", "Coverage"])
        self.assertIn("Projected Vacation Schedule", wb["Cover"]["A1"].value)


if __name__ == "__main__":
    unittest.main()
